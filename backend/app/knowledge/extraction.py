"""Upload validation, text extraction and chunking. All inputs are untrusted.

Validation is by content (magic bytes / structure), not by the client-supplied MIME type:
- PDF: starts with ``%PDF-``; page and character limits.
- DOCX: a ZIP containing ``word/document.xml``; entry-count and uncompressed-size limits
  (zip-bomb guard); macros are irrelevant because we only read text.
- TXT: UTF-8 (BOM allowed), no NUL bytes.
- Markdown (.md/.markdown): UTF-8 text; markup (headings, emphasis, links, images, code fences)
  is reduced to its readable text.
"""

import io
import re
import zipfile
from dataclasses import dataclass

MAX_PDF_PAGES = 300
MAX_TEXT_CHARS = 2_000_000
MAX_ZIP_ENTRIES = 2_000
MAX_ZIP_UNCOMPRESSED = 60 * 1024 * 1024
CHUNK_TARGET = 1_000
CHUNK_OVERLAP = 150

ALLOWED_EXTENSIONS = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ExtractionError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ExtractedDocument:
    kind: str  # pdf | docx | txt | md | markdown
    mime_type: str
    text: str


def _extension(filename: str) -> str:
    m = re.search(r"\.[a-z0-9]+$", filename.lower())
    return m.group(0) if m else ""


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("encrypted_pdf", "Password-protected PDFs are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ExtractionError("too_many_pages", f"PDFs are limited to {MAX_PDF_PAGES} pages")
        parts: list[str] = []
        total = 0
        for page in reader.pages:
            chunk = page.extract_text() or ""
            total += len(chunk)
            if total > MAX_TEXT_CHARS:
                raise ExtractionError("too_much_text", "Document contains too much text")
            parts.append(chunk)
        return "\n\n".join(parts)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise ExtractionError("unreadable_pdf", "The PDF could not be read") from exc


def _docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if (
                len(infos) > MAX_ZIP_ENTRIES
                or sum(i.file_size for i in infos) > MAX_ZIP_UNCOMPRESSED
            ):
                raise ExtractionError("docx_too_large", "The DOCX file expands to too much data")
            if "word/document.xml" not in zf.namelist():
                raise ExtractionError("not_docx", "The file is not a valid DOCX document")
    except zipfile.BadZipFile as exc:
        raise ExtractionError("not_docx", "The file is not a valid DOCX document") from exc
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises a variety of parser errors
        raise ExtractionError("unreadable_docx", "The DOCX file could not be read") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _txt(data: bytes) -> str:
    if b"\x00" in data:
        raise ExtractionError("binary_text", "Text files must not contain binary data")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("not_utf8", "Text files must be UTF-8 encoded") from exc


def _markdown(text: str) -> str:
    text = re.sub(r"^```[^\n]*$", "", text, flags=re.M)  # fence lines; code text is kept
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)  # images -> alt text
    text = re.sub(r"\[([^\]]+)\]\(([^)\s]+)[^)]*\)", r"\1 (\2)", text)  # links -> text (url)
    text = re.sub(r"^\s{0,3}#{1,6}\s+", "", text, flags=re.M)  # headings
    text = re.sub(r"^\s{0,3}>\s?", "", text, flags=re.M)  # block quotes
    text = re.sub(r"^\s*[-*+]\s+", "- ", text, flags=re.M)  # bullets
    text = re.sub(r"(\*\*|__|\*|_|`)(\S(?:.*?\S)?)\1", r"\2", text)  # emphasis / inline code
    text = re.sub(r"^\s*([-*_]\s*){3,}$", "", text, flags=re.M)  # horizontal rules
    return text


def extract(filename: str, data: bytes) -> ExtractedDocument:
    ext = _extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractionError(
            "unsupported_type", "Only PDF, DOCX, TXT and Markdown files are supported"
        )
    if ext == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ExtractionError("content_mismatch", "The file content is not a PDF")
        text = _pdf(data)
    elif ext == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ExtractionError("content_mismatch", "The file content is not a DOCX document")
        text = _docx(data)
    elif ext in (".md", ".markdown"):
        text = _markdown(_txt(data))
    else:
        text = _txt(data)
    text = _clean(text)
    if len(text) > MAX_TEXT_CHARS:
        raise ExtractionError("too_much_text", "Document contains too much text")
    if not text:
        raise ExtractionError("no_text", "No readable text was found in the document")
    return ExtractedDocument(kind=ext[1:], mime_type=ALLOWED_EXTENSIONS[ext], text=text)


def chunk_text(text: str, target: int = CHUNK_TARGET, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Paragraph-aware chunks of roughly ``target`` characters with a small overlap."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        while len(para) > target:  # split very long paragraphs on sentence boundaries
            cut = para.rfind(". ", 0, target)
            cut = cut + 1 if cut > target // 2 else target
            pieces = para[:cut].strip()
            para = para[cut:].strip()
            if current:
                chunks.append(current)
                current = ""
            chunks.append(pieces)
        if len(current) + len(para) + 2 <= target:
            current = f"{current}\n\n{para}" if current else para
        else:
            if current:
                chunks.append(current)
            tail = current[-overlap:] if current and overlap else ""
            current = f"{tail} {para}".strip() if tail else para
    if current:
        chunks.append(current)
    return chunks
